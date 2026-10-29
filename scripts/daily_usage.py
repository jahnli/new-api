#!/usr/bin/env python3
"""
按人员统计 2026 年指定月份每天的人民币消费金额。

数据口径：
  - 金额依据 quota_data.quota 计算。
  - 人民币金额 = quota / QuotaPerUnit * 6.8。
  - 仅统计 users.open_id 或 users.dingding_id 非空的人员。
  - 相同工号的多个用户合并为一行；无工号用户按 user_id 分开统计。
  - users 表优先使用 job_number 字段，当前数据库不存在该字段时自动使用
    employee_number 字段；导出列名统一为“工号”。
  - 周六、周日金额固定归入下一个周一，按归并后的日期所属月份统计。
  - 默认分别导出赛美特集团、上海繁竞科技有限公司，各公司独立合并工号。
  - 输出文件名自动附加公司名称，--output 用于指定基础文件名。

示例：
  python daily_usage.py --month 6
  python daily_usage.py --month 6 --company 赛美特集团
  python daily_usage.py --month 7 --output 2026年7月人员每日金额.xlsx
  python daily_usage.py --month 7 --output 2026_07_daily_amount.csv
"""

from __future__ import annotations

import argparse
import calendar
import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

YEAR = 2026

POSTGRES_HOST = "10.16.7.77"
POSTGRES_PORT = 5432
POSTGRES_DATABASE = "ai-gateway"
POSTGRES_USER = "admin"
POSTGRES_PASSWORD = "semitech"
POSTGRES_SCHEMA = "public"
DEFAULT_QUOTA_PER_UNIT = 500000.0
DEFAULT_EXCHANGE_RATE = 6.8
TARGET_COMPANIES = ("赛美特集团", "上海繁竞科技有限公司")


@dataclass
class PersonUsage:
    representative_user_id: int
    display_name: str
    username: str
    job_number: str
    user_ids: set[int] = field(default_factory=set)
    daily_quota: dict[int, int] = field(default_factory=dict)
    department_name: str = ""


class DB:
    def __init__(self, connection: Any):
        self.connection = connection

    def query(self, sql: str, params: tuple[Any, ...] = ()) -> list[dict[str, Any]]:
        cursor = self.connection.cursor()
        try:
            cursor.execute(sql, params)
            columns = [description[0] for description in cursor.description]
            return [dict(zip(columns, row)) for row in cursor.fetchall()]
        finally:
            cursor.close()

    def close(self) -> None:
        self.connection.close()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="从 quota_data 统计 2026 年指定月份每个人每天的人民币消费金额"
    )
    parser.add_argument("--month", type=int, required=True, help="统计月份，范围 1-12")
    parser.add_argument("--tz", default="Asia/Shanghai", help="统计时区，默认 Asia/Shanghai")
    parser.add_argument(
        "--output", help="输出基础文件名，支持 .xlsx 或 .csv；自动附加公司名称"
    )
    parser.add_argument(
        "--company",
        action="append",
        choices=TARGET_COMPANIES,
        help="要统计的公司，可重复指定；默认分别统计两家公司",
    )
    parser.add_argument(
        "--quota-per-unit",
        type=float,
        help="额度换算金额的分母；默认读取 options.QuotaPerUnit，读取不到使用 500000",
    )
    parser.add_argument(
        "--exchange-rate",
        type=float,
        default=DEFAULT_EXCHANGE_RATE,
        help="1 USD 对应人民币汇率，默认 6.8",
    )
    parser.add_argument("--host", default=POSTGRES_HOST, help="PostgreSQL host")
    parser.add_argument("--port", type=int, default=POSTGRES_PORT, help="PostgreSQL port")
    parser.add_argument("--user", default=POSTGRES_USER, help="PostgreSQL user")
    parser.add_argument("--password", default=POSTGRES_PASSWORD, help="PostgreSQL password")
    parser.add_argument("--schema", default=POSTGRES_SCHEMA, help="PostgreSQL schema")
    args = parser.parse_args()

    if not 1 <= args.month <= 12:
        parser.error("--month 必须在 1 到 12 之间")
    if args.quota_per_unit is not None and args.quota_per_unit <= 0:
        parser.error("--quota-per-unit 必须大于 0")
    if args.exchange_rate <= 0:
        parser.error("--exchange-rate 必须大于 0")
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", args.schema):
        parser.error("--schema 只能包含字母、数字和下划线，且不能以数字开头")
    args.companies = list(dict.fromkeys(args.company or TARGET_COMPANIES))
    return args


def connect_postgres(args: argparse.Namespace) -> DB:
    try:
        import psycopg2
    except ImportError as exc:
        raise RuntimeError("缺少 psycopg2，请先运行：python -m pip install psycopg2-binary") from exc

    connection = psycopg2.connect(
        host=args.host,
        port=args.port,
        dbname=POSTGRES_DATABASE,
        user=args.user,
        password=args.password,
    )
    cursor = connection.cursor()
    try:
        cursor.execute(f'SET search_path TO "{args.schema}"')
    finally:
        cursor.close()
    return DB(connection)


def month_range(month: int, tz: ZoneInfo) -> tuple[datetime, datetime]:
    start = datetime.combine(date(YEAR, month, 1), time.min, tzinfo=tz)
    if month == 12:
        next_month = date(YEAR + 1, 1, 1)
    else:
        next_month = date(YEAR, month + 1, 1)
    end = datetime.combine(next_month, time.min, tzinfo=tz)
    return start, end


def find_job_number_column(db: DB, schema: str) -> str:
    rows = db.query(
        """
        SELECT column_name
        FROM information_schema.columns
        WHERE table_schema = %s
          AND table_name = 'users'
          AND column_name IN ('job_number', 'employee_number')
        """,
        (schema,),
    )
    available = {str(row["column_name"]) for row in rows}
    if "job_number" in available:
        return "job_number"
    if "employee_number" in available:
        return "employee_number"
    raise RuntimeError("users 表中既没有 job_number，也没有 employee_number 字段")


def fetch_quota_per_unit(db: DB, configured_value: float | None) -> float:
    if configured_value is not None:
        return configured_value
    rows = db.query("SELECT value FROM options WHERE key = %s LIMIT 1", ("QuotaPerUnit",))
    if not rows:
        return DEFAULT_QUOTA_PER_UNIT
    try:
        value = float(rows[0].get("value") or 0)
    except (TypeError, ValueError):
        return DEFAULT_QUOTA_PER_UNIT
    return value if value > 0 else DEFAULT_QUOTA_PER_UNIT


def fetch_eligible_users(
    db: DB,
    job_column: str,
    company: str | None = None,
) -> list[dict[str, Any]]:
    # job_column 只能来自 find_job_number_column 中的固定白名单。
    company_clause = ""
    params: tuple[Any, ...] = ()
    if company is not None:
        company_clause = "AND company = %s"
        params = (company,)
    return db.query(
        f"""
        SELECT
          id,
          COALESCE(display_name, '') AS display_name,
          COALESCE(username, '') AS username,
          COALESCE({job_column}, '') AS job_number,
          COALESCE(department_name, '') AS department_name
        FROM users
        WHERE (BTRIM(COALESCE(open_id, '')) <> ''
          OR BTRIM(COALESCE(dingding_id, '')) <> '')
          {company_clause}
        ORDER BY id ASC
        """,
        params,
    )


def build_people(users: list[dict[str, Any]]) -> tuple[list[PersonUsage], dict[int, PersonUsage]]:
    people_by_key: dict[tuple[str, str | int], PersonUsage] = {}
    person_by_user_id: dict[int, PersonUsage] = {}

    for row in users:
        user_id = int(row["id"])
        job_number = str(row.get("job_number") or "").strip()
        # 空工号不能作为共享合并键，否则所有空工号人员会被合并成一人。
        key: tuple[str, str | int] = ("job", job_number) if job_number else ("user", user_id)
        person = people_by_key.get(key)
        if person is None:
            person = PersonUsage(
                representative_user_id=user_id,
                display_name=str(row.get("display_name") or ""),
                username=str(row.get("username") or ""),
                job_number=job_number,
                department_name=str(row.get("department_name") or ""),
            )
            people_by_key[key] = person
        person.user_ids.add(user_id)
        person_by_user_id[user_id] = person

    people = sorted(
        people_by_key.values(),
        key=lambda person: (
            person.job_number == "",
            person.job_number.casefold(),
            person.display_name.casefold(),
            person.representative_user_id,
        ),
    )
    return people, person_by_user_id


def fetch_daily_usage(
    db: DB,
    start_ts: int,
    end_ts: int,
    timezone_name: str,
) -> list[dict[str, Any]]:
    return db.query(
        """
        SELECT
          user_id,
          (TO_TIMESTAMP(created_at) AT TIME ZONE %s)::date AS usage_date,
          COALESCE(SUM(quota), 0) AS quota
        FROM quota_data
        WHERE created_at >= %s
          AND created_at < %s
        GROUP BY user_id, usage_date
        ORDER BY user_id ASC, usage_date ASC
        """,
        (timezone_name, start_ts, end_ts),
    )


def apply_daily_usage(
    rows: list[dict[str, Any]],
    person_by_user_id: dict[int, PersonUsage],
    target_year: int | None = None,
    target_month: int | None = None,
) -> None:
    for row in rows:
        person = person_by_user_id.get(int(row.get("user_id") or 0))
        if person is None:
            continue
        usage_date = row.get("usage_date")
        if isinstance(usage_date, datetime):
            usage_date = usage_date.date()
        if not isinstance(usage_date, date):
            continue
        if usage_date.weekday() >= 5:
            usage_date += timedelta(days=7 - usage_date.weekday())
        if target_year is not None and usage_date.year != target_year:
            continue
        if target_month is not None and usage_date.month != target_month:
            continue
        day = usage_date.day
        person.daily_quota[day] = person.daily_quota.get(day, 0) + int(row.get("quota") or 0)


def amount_cny(
    quota: int,
    quota_per_unit: float,
    exchange_rate: float,
    decimals: int = 2,
) -> float:
    return round(quota / quota_per_unit * exchange_rate, decimals)


def make_table(
    people: list[PersonUsage],
    month: int,
    quota_per_unit: float,
    exchange_rate: float,
    include_department: bool = False,
    amount_decimals: int = 2,
    day_column_count: int | None = None,
) -> tuple[list[str], list[list[Any]]]:
    days_in_month = calendar.monthrange(YEAR, month)[1]
    output_days = day_column_count if day_column_count is not None else days_in_month
    if output_days < days_in_month or output_days > 31:
        raise ValueError("日期列数不能少于当月实际天数，且不能超过 31")
    identity_headers = [
        "年份",
        "月份",
        "员工名称",
        "username",
        "工号",
    ]
    if include_department:
        identity_headers.append("部门")
    headers = [
        *identity_headers,
        *[f"{day}日" for day in range(1, output_days + 1)],
        "总计",
    ]

    rows: list[list[Any]] = []
    for person in people:
        daily_values = [
            amount_cny(
                person.daily_quota.get(day, 0),
                quota_per_unit,
                exchange_rate,
                decimals=amount_decimals,
            )
            for day in range(1, output_days + 1)
        ]
        identity_values: list[Any] = [
            YEAR,
            month,
            person.display_name,
            person.username,
            person.job_number,
        ]
        if include_department:
            identity_values.append(person.department_name)
        rows.append([
            *identity_values,
            *daily_values,
            round(sum(daily_values), amount_decimals),
        ])
    return headers, rows


def write_csv(output: Path, headers: list[str], rows: list[list[Any]]) -> None:
    with output.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)


def write_xlsx(
    output: Path,
    headers: list[str],
    rows: list[list[Any]],
    month: int,
    amount_decimals: int = 2,
) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("输出 .xlsx 需要 openpyxl：python -m pip install openpyxl") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = f"{YEAR}年{month}月每日金额"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    first_daily_column = headers.index("1日") + 1
    sheet.freeze_panes = f"{get_column_letter(first_daily_column)}2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    widths = {
        "年份": 9,
        "月份": 7,
        "员工名称": 24,
        "username": 22,
        "工号": 20,
        "部门": 42,
    }
    for column_index in range(1, len(headers) + 1):
        header = headers[column_index - 1]
        sheet.column_dimensions[get_column_letter(column_index)].width = widths.get(header, 14)
    for row in sheet.iter_rows(min_row=2, min_col=first_daily_column, max_col=len(headers)):
        for cell in row:
            cell.number_format = f"#,##0.{('0' * amount_decimals)}"

    workbook.save(output)


def resolve_output(args: argparse.Namespace, company: str) -> Path:
    output = (
        Path(args.output)
        if args.output
        else Path(f"person_daily_amount_{YEAR}{args.month:02d}_weekend_to_monday.xlsx")
    )
    if output.suffix.lower() not in {".xlsx", ".csv"}:
        raise ValueError("输出文件只支持 .xlsx 或 .csv")
    return output.with_name(f"{output.stem}_{company}{output.suffix}")


def main() -> None:
    args = parse_args()
    timezone = ZoneInfo(args.tz)
    start, end = month_range(args.month, timezone)
    outputs = {company: resolve_output(args, company) for company in args.companies}
    # 月初为周日或周一时，补取上月周末，归入本月第一个周一。
    fetch_start = start
    if start.weekday() == 0:
        fetch_start = start - timedelta(days=2)
    elif start.weekday() == 6:
        fetch_start = start - timedelta(days=1)

    db = connect_postgres(args)
    try:
        job_column = find_job_number_column(db, args.schema)
        quota_per_unit = fetch_quota_per_unit(db, args.quota_per_unit)
        users_by_company = {
            company: fetch_eligible_users(db, job_column, company=company)
            for company in args.companies
        }
        usage_rows = fetch_daily_usage(
            db,
            int(fetch_start.timestamp()),
            int(end.timestamp()),
            args.tz,
        )
    finally:
        db.close()

    print(f"归集月份: {YEAR}年{args.month}月（按归并后的日期筛选）")
    print(f"数据查询区间: {fetch_start:%Y-%m-%d %H:%M:%S %Z} <= created_at < {end:%Y-%m-%d %H:%M:%S %Z}")
    print(f"工号字段: users.{job_column}")
    print(f"金额公式: quota / {quota_per_unit:g} * {args.exchange_rate:g}")
    print("周末金额归入下一个周一；跨月周末计入周一所属月份")

    for company, users in users_by_company.items():
        people, person_by_user_id = build_people(users)
        apply_daily_usage(
            usage_rows,
            person_by_user_id,
            target_year=YEAR,
            target_month=args.month,
        )
        headers, table_rows = make_table(
            people,
            args.month,
            quota_per_unit,
            args.exchange_rate,
            include_department=True,
            amount_decimals=2,
            day_column_count=31,
        )
        output = outputs[company]
        output.parent.mkdir(parents=True, exist_ok=True)
        if output.suffix.lower() == ".csv":
            write_csv(output, headers, table_rows)
        else:
            write_xlsx(output, headers, table_rows, args.month, amount_decimals=2)

        merged_groups = sum(1 for person in people if len(person.user_ids) > 1)
        total_amount = round(sum(float(row[-1]) for row in table_rows), 2)
        print(f"\n公司筛选: users.company = {company}")
        print(f"open_id 或 dingding_id 非空用户数: {len(users)}")
        print(f"合并后人员数: {len(people)}（其中 {merged_groups} 个工号包含多个用户）")
        print(f"总金额(人民币): {total_amount:,.2f}")
        print(f"已输出: {output.resolve()}")


if __name__ == "__main__":
    main()
